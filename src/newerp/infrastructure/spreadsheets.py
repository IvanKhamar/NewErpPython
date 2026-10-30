"""Excel .xlsx files through openpyxl (ClosedXML in the .NET version)."""

import io
import zipfile
from collections.abc import Sequence
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from newerp.domain.common import DomainError, plain


class OpenpyxlSpreadsheet:
    def write(self, sheet_name: str, header: Sequence[str], rows: Sequence[Sequence[object]]) -> bytes:
        book = Workbook()
        sheet = book.active
        sheet.title = sheet_name
        sheet.append(list(header))
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            # Text stays text: the phone "0671234567" must not become a number without its zero.
            sheet.append([float(v) if isinstance(v, Decimal) else v for v in row])
        sheet.freeze_panes = "A2"

        for column, cells in enumerate(sheet.columns, start=1):
            width = max((len(str(c.value)) for c in cells if c.value is not None), default=0)
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 1), 100)

        buffer = io.BytesIO()
        book.save(buffer)
        return buffer.getvalue()

    def read(self, content: bytes) -> list[list[str]]:
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as error:
            raise DomainError("Не вдалося прочитати файл. Потрібен файл Excel (.xlsx).") from error
        try:
            sheet = book.worksheets[0]
            rows = [[_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        finally:
            book.close()
        # Like ClosedXML's RangeUsed: trailing empty rows and columns do not count.
        while rows and all(cell == "" for cell in rows[-1]):
            rows.pop()
        width = max((max((i + 1 for i, cell in enumerate(row) if cell != ""), default=0) for row in rows), default=0)
        return [(row + [""] * width)[:width] for row in rows]


def _text(value: object) -> str:
    """A cell as the user sees it: 12.5 rather than 12.500000001, whole numbers without ".0"."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return plain(Decimal(repr(value)))
    if isinstance(value, int):
        return str(value)
    if isinstance(value, datetime):
        return (
            value.strftime("%d.%m.%Y %H:%M:%S") if value.time() != datetime.min.time() else value.strftime("%d.%m.%Y")
        )
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    return str(value)
